"""Builds the BCSE497J project report in the department template's format.

FORMATTING is taken from the specification written inside the template document:

    Page size      A4
    Left margin    3.81 cm      Right / Top / Bottom margin   2.54 cm
    Body font      Times New Roman 12 pt, justified, 1.5 line spacing
    Chapter head   Times New Roman 16 pt, Bold, CAPS
    Section head   Times New Roman 14 pt, CAPS
    Subsection     Times New Roman 12 pt, CAPS
    Page numbers   Arabic numerals, bottom centre
    References     APA

Note that the template .docx file's own section properties disagree with the
specification it prints (its instruction pages use Letter paper and ~0.7 inch
margins). The written specification is followed here, since that is what the
department is asking for.

Every numeric result is read from results/*.json so the report cannot quote a
figure the project does not produce. Details the author must supply -- names,
registration numbers, dates -- are left as << >> placeholders.
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from src.config import FIGURES_DIR, PROJECT_ROOT, RESULTS_DIR

OUT = os.path.join(PROJECT_ROOT, 'BCSE497J_Project_Report.docx')

FONT = 'Times New Roman'
BODY_PT = 12
CHAPTER_PT = 16
SECTION_PT = 14
SUBSECTION_PT = 12
LINE_SPACING = 1.5

PH = '<<{}>>'          # placeholder marker


# --------------------------------------------------------------- infrastructure
def set_cell_font(cell, size=BODY_PT, bold=False, align=None):
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(size)
            r.bold = bold


def add_page_number_footer(section):
    """Arabic page number, bottom centre, as the specification requires."""
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run()
    run.font.name = FONT
    run.font.size = Pt(BODY_PT)
    for instr, kind in (('begin', 'fldChar'), ('PAGE', 'instrText'), ('end', 'fldChar')):
        el = OxmlElement(f'w:{kind}')
        if kind == 'fldChar':
            el.set(qn('w:fldCharType'), instr)
        else:
            el.set(qn('xml:space'), 'preserve')
            el.text = ' PAGE '
        run._r.append(el)


class Report:
    def __init__(self):
        self.doc = Document()
        self._page_setup()
        self._base_style()
        self.figures = []       # (number, caption)
        self.tables = []        # (number, caption)

    def _page_setup(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)        # A4
        s.left_margin = Cm(3.81)
        s.right_margin = Cm(2.54)
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.header_distance = Cm(1.25)
        s.footer_distance = Cm(1.25)
        add_page_number_footer(s)

    def _base_style(self):
        st = self.doc.styles['Normal']
        st.font.name = FONT
        st.font.size = Pt(BODY_PT)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        pf = st.paragraph_format
        pf.line_spacing = LINE_SPACING
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # ----------------------------------------------------------- primitives
    def para(self, text='', *, size=BODY_PT, bold=False, italic=False, caps=False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=LINE_SPACING, before=0, after=0,
             color=None, indent=None, underline=False):
        p = self.doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.line_spacing = spacing
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        if indent is not None:
            pf.left_indent = Cm(indent)
        if text:
            r = p.add_run(text.upper() if caps else text)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
            r.underline = underline
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    def blank(self, n=1, size=BODY_PT):
        for _ in range(n):
            self.para('', size=size, spacing=1.0)

    def page_break(self):
        self.doc.add_page_break()

    # ----------------------------------------------------------- headings
    def chapter(self, number, title):
        """Chapter Heading -- 16 pt, Bold, Times New Roman."""
        self.page_break()
        self.para(f'CHAPTER {number}', size=CHAPTER_PT, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, after=6)
        self.para(title.upper(), size=CHAPTER_PT, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, after=18)

    def section(self, number, title):
        """Section Heading -- 14 pt, Times New Roman, CAPS."""
        self.para(f'{number}  {title.upper()}', size=SECTION_PT,
                  align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.5, before=12, after=6)

    def subsection(self, number, title):
        """Subsection Heading -- 12 pt, Times New Roman, CAPS."""
        self.para(f'{number}  {title.upper()}', size=SUBSECTION_PT,
                  align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.5, before=10, after=4)

    def unnumbered_heading(self, title):
        self.para(title.upper(), size=CHAPTER_PT, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, after=18)

    # ----------------------------------------------------------- content
    def bullets(self, items, indent=1.0):
        for it in items:
            p = self.doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(indent)
            p.paragraph_format.line_spacing = LINE_SPACING
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(it)
            r.font.name = FONT
            r.font.size = Pt(BODY_PT)

    def numbered(self, items, indent=1.0):
        for it in items:
            p = self.doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Cm(indent)
            p.paragraph_format.line_spacing = LINE_SPACING
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(it)
            r.font.name = FONT
            r.font.size = Pt(BODY_PT)

    def table(self, number, caption, headers, rows, widths=None, font=11):
        self.tables.append((number, caption))
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = str(h)
            set_cell_font(c, size=font, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                set_cell_font(cells[i], size=font,
                              align=WD_ALIGN_PARAGRAPH.CENTER if i else
                              WD_ALIGN_PARAGRAPH.LEFT)
        if widths:
            for r in t.rows:
                for i, w in enumerate(widths):
                    r.cells[i].width = Cm(w)
        self.para(f'Table {number}  {caption}', size=11,
                  align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=4, after=10)
        return t

    def figure(self, number, caption, filename, width=14.0, subdir=None):
        base = os.path.join(FIGURES_DIR, subdir) if subdir else FIGURES_DIR
        path = os.path.join(base, filename)
        self.figures.append((number, caption))
        if os.path.exists(path):
            self.doc.add_picture(path, width=Cm(width))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.doc.paragraphs[-1].paragraph_format.space_before = Pt(8)
        else:
            self.para(f'[ FIGURE {number} TO BE INSERTED: {filename} ]',
                      align=WD_ALIGN_PARAGRAPH.CENTER, color='FF0000', before=8)
        self.para(f'Figure {number}  {caption}', size=11,
                  align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=4, after=12)

    def placeholder_figure(self, number, caption, note):
        self.figures.append((number, caption))
        self.para(f'[ INSERT FIGURE {number} HERE — {note} ]',
                  align=WD_ALIGN_PARAGRAPH.CENTER, color='FF0000', before=10, after=4)
        self.para(f'Figure {number}  {caption}', size=11,
                  align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, after=12)

    def save(self):
        self.doc.save(OUT)
        return OUT


# --------------------------------------------------------------- project data
def load_results():
    r = {}
    for name in ('phase2', 'phase3', 'phase4', 'phase5', 'phase5b', 'phase6_final'):
        with open(os.path.join(RESULTS_DIR, f'{name}_results.json'), encoding='utf-8') as fh:
            r[name] = json.load(fh)
    for tag in ('base', 'v3', 'v3_ens3'):
        with open(os.path.join(RESULTS_DIR, f'phase8_kfold_cv_{tag}_results.json'),
                  encoding='utf-8') as fh:
            r[f'cv_{tag}'] = json.load(fh)['summary']
    with open(os.path.join(RESULTS_DIR, 'phase6_search_results.json'), encoding='utf-8') as fh:
        r['search'] = json.load(fh)
    with open(os.path.join(RESULTS_DIR, 'phase7_ablations_results.json'),
              encoding='utf-8') as fh:
        r['p7'] = json.load(fh)
    with open(os.path.join(RESULTS_DIR, 'ablation_encoder_study_results.json'),
              encoding='utf-8') as fh:
        r['abl'] = json.load(fh)
    return r


def p7(r, needle):
    for e in r['p7'].values():
        if isinstance(e, dict) and needle in e.get('label', ''):
            return e['test_metrics']
    raise KeyError(needle)


def abl_stat(r, key, metric):
    import numpy as np
    v = [x[metric] for x in r['abl'].values() if x['key'] == key]
    return float(np.mean(v)), float(np.std(v))


def cv(r, tag, metric):
    m = r[f'cv_{tag}'][metric]
    return m['mean'], m['std']


# --------------------------------------------------------------- front matter
TITLE = ('LESION-GUIDED MULTI-TASK DEEP LEARNING FOR CLASSIFICATION AND '
         'SEGMENTATION OF KIDNEY DISEASES IN ULTRASOUND IMAGES')
C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY
L = WD_ALIGN_PARAGRAPH.LEFT


def title_page(rp, cover=False):
    """Title / cover page, following the template's per-line font specification."""
    rp.blank(2)
    rp.para('A project report on', size=12, italic=True, align=C, spacing=1.0, after=18)
    rp.blank(1)
    rp.para(TITLE, size=20, bold=True, align=C, spacing=1.25, after=18)
    rp.blank(1)
    rp.para('Submitted in partial fulfillment for the award of the degree of',
            size=14, italic=True, align=C, spacing=1.0, after=14)
    rp.para(f'Bachelor of Technology in {PH.format("PROGRAM NAME WITH SPECIALIZATION")}',
            size=22, bold=True, align=C, spacing=1.15, after=16)
    rp.blank(1)
    rp.para('by', size=14, italic=True, align=C, spacing=1.0, after=12)
    rp.para(f'{PH.format("NAME OF THE CANDIDATE")}  ({PH.format("REG. NO.")})',
            size=16, bold=True, align=C, spacing=1.15, after=4)
    rp.para(f'{PH.format("SECOND MEMBER NAME")}  ({PH.format("REG. NO.")})',
            size=16, bold=True, align=C, spacing=1.15, after=4)
    rp.para(f'{PH.format("THIRD MEMBER NAME")}  ({PH.format("REG. NO.")})',
            size=16, bold=True, align=C, spacing=1.15, after=14)
    rp.blank(1)
    rp.para('[ INSERT VIT LOGO HERE ]', size=12, align=C, color='FF0000',
            spacing=1.0, after=14)
    rp.blank(1)
    rp.para('SCHOOL OF COMPUTER SCIENCE AND ENGINEERING', size=16, bold=True,
            align=C, spacing=1.15, after=10)
    rp.para(f'{PH.format("MONTH")}, {PH.format("YEAR")}', size=12, align=C,
            spacing=1.0)
    rp.page_break()


def declaration(rp):
    rp.blank(1)
    rp.para('DECLARATION', size=14, bold=True, underline=True, align=C,
            spacing=1.5, after=18)
    rp.blank(1)
    rp.para(f'I hereby declare that the thesis entitled \u201c{TITLE.title()}\u201d '
            f'submitted by {PH.format("NAME")} ({PH.format("REGISTER NUMBER")}), '
            f'for the award of the degree of Bachelor of Technology in '
            f'{PH.format("PROGRAM NAME WITH SPECIALIZATION")}, Vellore Institute of '
            f'Technology, Chennai, is a record of bonafide work carried out by me '
            f'under the supervision of {PH.format("GUIDE NAME")}.',
            align=J, after=12)
    rp.para('I further declare that the work reported in this thesis has not been '
            'submitted and will not be submitted, either in part or in full, for the '
            'award of any other degree or diploma in this institute or any other '
            'institute or university.', align=J, after=12)
    rp.blank(3)
    rp.para('Place: Chennai', align=L, spacing=1.5, after=6)
    rp.para(f'Date: {PH.format("DATE")}', align=L, spacing=1.5, after=24)
    rp.blank(2)
    rp.para('Signature of the Candidate', size=12, bold=True,
            align=WD_ALIGN_PARAGRAPH.RIGHT, spacing=1.0)
    rp.page_break()


def certificate(rp):
    rp.para('School of Computer Science and Engineering', size=14, bold=True,
            align=C, spacing=1.5, after=14)
    rp.para('CERTIFICATE', size=17, bold=True, align=C, spacing=1.5, after=18)
    rp.para(f'This is to certify that the report entitled \u201c{TITLE.title()}\u201d is '
            f'prepared and submitted by {PH.format("NAME")} '
            f'({PH.format("REGISTER NUMBER")}) to Vellore Institute of Technology, '
            f'Chennai, in partial fulfillment of the requirement for the award of the '
            f'degree of Bachelor of Technology in '
            f'{PH.format("PROGRAM NAME WITH SPECIALIZATION")} is a bonafide record '
            f'carried out under my guidance. The project fulfills the requirements as '
            f'per the regulations of this University and in my opinion meets the '
            f'necessary standards for submission. The contents of this report have not '
            f'been submitted and will not be submitted either in part or in full, for '
            f'the award of any other degree or diploma and the same is certified.',
            align=J, after=24)
    rp.para('Signature of the Guide:', align=L, spacing=1.5, after=10)
    rp.para(f'Name: Dr./Prof. {PH.format("GUIDE NAME")}', align=L, spacing=1.5, after=10)
    rp.para(f'Date: {PH.format("DATE")}', align=L, spacing=1.5, after=24)
    rp.blank(2)

    t = rp.doc.add_table(rows=3, cols=2)
    data = [('Signature of the Examiner', 'Signature of the Examiner'),
            (f'Name: {PH.format("EXAMINER 1")}', f'Name: {PH.format("EXAMINER 2")}'),
            (f'Date: {PH.format("DATE")}', f'Date: {PH.format("DATE")}')]
    for i, (a, b) in enumerate(data):
        t.rows[i].cells[0].text = a
        t.rows[i].cells[1].text = b
        set_cell_font(t.rows[i].cells[0], align=L)
        set_cell_font(t.rows[i].cells[1], align=L)
    rp.blank(2)

    rp.para('Approved by the Head of Department,', align=C, spacing=1.5, after=4)
    rp.para(PH.format('PROGRAM NAME WITH SPECIALIZATION'), size=12, bold=True,
            align=C, spacing=1.5, after=14)
    rp.para(f'Name: {PH.format("HoD NAME")}', align=L, spacing=1.5, after=10)
    rp.para(f'Date: {PH.format("DATE")}', align=L, spacing=1.5, after=18)
    rp.para('(Seal of the School)', size=10, bold=True, align=L, spacing=1.0)
    rp.page_break()
