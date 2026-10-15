"""Builds the lesion-guided pooling diagram as a PPTX slide.

Styling is matched exactly to the user's decoder_diagrams.pptx so the new slide
drops into that deck without looking out of place:

    slide       13.33 x 7.50 in, TITLE_ONLY layout, title 26 pt bold
    panel label Calibri 14 pt bold  #222222
    box title   Calibri 12 pt bold  #222222
    box detail  Calibri 10 pt       #666666
    row box     height 0.45 in, pitch 0.625 in
    connectors  #555555 at 1.25 pt, green path #2E8B3E
    palette     tensor  FCE9D4 / D98A2B      module  E4E4F7 / 5B5FC7
                loss    FBE3E3 / C0392B      vector  E2F4E2 / 2E8B3E
                output  E3EEFB / 3A78C2      stop    FBE3E3 / C0392B (heavy)

Content follows model.py MultiTaskNetV3.forward and losses.compute_loss.
"""
import copy
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from src.config import PROJECT_ROOT

OUT = os.path.join(PROJECT_ROOT, 'paper', 'pooling_diagram.pptx')

# ---------------------------------------------------------------- palette
TENSOR = ('FCE9D4', 'D98A2B')
MODULE = ('E4E4F7', '5B5FC7')
LOSSC = ('FBE3E3', 'C0392B')
VECTOR = ('E2F4E2', '2E8B3E')
OUTPUT = ('E3EEFB', '3A78C2')

INK, MUTED, SHAPE_TXT = '222222', '666666', '333333'
WIRE, GREEN = '555555', '2E8B3E'

# ---------------------------------------------------------------- geometry
ROW_H = 0.45
PITCH = 0.625
TOP0 = 1.22

GCOL_L, GCOL_W = 0.40, 2.00          # global-average-pooling branch
MAIN_L, MAIN_W = 2.85, 4.30          # main chain
RPAN_L, RPAN_W = 7.55, 5.50          # notes / loss / legend
MAIN_CX = MAIN_L + MAIN_W / 2
GCOL_CX = GCOL_L + GCOL_W / 2


def row_top(i):
    return TOP0 + i * PITCH


def row_mid(i):
    return row_top(i) + ROW_H / 2


# ---------------------------------------------------------------- helpers
def arrow(shapes, x1, y1, x2, y2, color=WIRE, width=1.25, head=True):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                             Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(width)
    if head:
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle',
                                                'w': 'med', 'len': 'med'})
        ln.append(tail)
    return c


def wire(shapes, pts, color=WIRE, width=1.25):
    """Orthogonal polyline as a sequence of straight segments; head on the last."""
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        arrow(shapes, x1, y1, x2, y2, color=color, width=width,
              head=(i == len(pts) - 2))


def box(shapes, left, top, width, height, fill, line, title, detail=None,
        title_pt=12, detail_pt=10, heavy=False):
    sh = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(left), Inches(top), Inches(width), Inches(height))
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.color.rgb = RGBColor.from_string(line)
    sh.line.width = Pt(2.0 if heavy else 1.0)
    sh.shadow.inherit = False

    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name = 'Calibri'
    r.font.size = Pt(title_pt)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(line if heavy else INK)

    if detail:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = detail
        r2.font.name = 'Calibri'
        r2.font.size = Pt(detail_pt)
        r2.font.color.rgb = RGBColor.from_string(MUTED)
    return sh


def label(shapes, left, top, width, height, text, pt=11, color=SHAPE_TXT,
          bold=False, align=PP_ALIGN.LEFT):
    tb = shapes.add_textbox(Inches(left), Inches(top), Inches(width),
                            Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    first = True
    for line_txt in text.split('\n'):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        r = p.add_run()
        r.text = line_txt
        r.font.name = 'Calibri'
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


# ---------------------------------------------------------------- build
def build():
    prs = Presentation()
    # exact EMU of the existing deck, so the slide pastes in without rescaling
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    layout = next(l for l in prs.slide_layouts if l.name == 'Title Only')
    slide = prs.slides.add_slide(layout)
    sh = slide.shapes

    title = slide.shapes.title
    title.left, title.top = Inches(0.50), Inches(0.30)
    title.width, title.height = Inches(12.33), Inches(0.55)
    tp = title.text_frame.paragraphs[0]
    tr = tp.add_run()
    tr.text = 'Lesion-guided pooling: how the classifier is given localised evidence'
    tr.font.size = Pt(26)
    tr.font.bold = True

    label(sh, MAIN_L, 0.92, 3.0, 0.30, 'Lesion-guided pooling path', pt=14,
          color=INK, bold=True)
    label(sh, GCOL_L, 0.92, 2.2, 0.30, 'Global branch', pt=14, color=INK,
          bold=True)
    label(sh, RPAN_L, 0.92, 3.0, 0.30, 'Supervision and notes', pt=14, color=INK,
          bold=True)

    # ---------------- main chain ------------------------------------------
    rows = [
        (TENSOR, 'Encoder feature map  F', '512 x 16 x 16', False),
        (MODULE, 'Lesion decoder', '5 up-blocks + conv 1x1', False),
        (MODULE, 'Lesion logits  L', '512 x 512', False),
        (MODULE, 'Sigmoid of L', 'predicted lesion probability map', False),
        (LOSSC, 'STOP-GRADIENT', 'detached from the computation graph', True),
        (OUTPUT, 'Resample to 16 x 16, normalise to unit sum',
         'gives the attention map  A', False),
        (VECTOR, 'Weighted sum over space', 'lesion vector = sum of A x F   (512-d)',
         False),
        (VECTOR, 'Concatenate  [ global vector || lesion vector ]', '1024-d', False),
        (OUTPUT, 'Local-pathology head', 'linear 1024 -> 9 logits', False),
    ]
    # Arrows leaving the stop-gradient carry the attention map only, which is
    # detached -- drawn green so the legend means something. The arrow out of
    # the weighted sum is grey again, because the lesion vector still passes
    # gradient to the encoder through F; only A is cut.
    DETACHED = {5, 6}
    for i, (pal, t, d, heavy) in enumerate(rows):
        box(sh, MAIN_L, row_top(i), MAIN_W, ROW_H, pal[0], pal[1], t, d,
            heavy=heavy)
        if i:
            arrow(sh, MAIN_CX, row_top(i - 1) + ROW_H, MAIN_CX, row_top(i),
                  color=GREEN if i in DETACHED else WIRE)

    # ---------------- global branch ---------------------------------------
    box(sh, GCOL_L, row_top(5), GCOL_W, ROW_H, VECTOR[0], VECTOR[1],
        'Global average pooling', 'over all 16 x 16 positions')
    box(sh, GCOL_L, row_top(6), GCOL_W, ROW_H, VECTOR[0], VECTOR[1],
        'Global vector', '512-d')
    arrow(sh, GCOL_CX, row_top(5) + ROW_H, GCOL_CX, row_top(6))

    # F -> global average pooling
    wire(sh, [(MAIN_L, row_mid(0)), (GCOL_CX, row_mid(0)),
              (GCOL_CX, row_top(5))])
    # global vector -> concatenation
    wire(sh, [(GCOL_CX, row_top(6) + ROW_H), (GCOL_CX, row_mid(7)),
              (MAIN_L, row_mid(7))])

    # ---------------- right panel -----------------------------------------
    box(sh, RPAN_L, row_top(2), RPAN_W, ROW_H, LOSSC[0], LOSSC[1],
        'Soft Dice loss  (sole supervision of L)',
        'the lesion decoder is trained by this term alone')
    arrow(sh, MAIN_L + MAIN_W, row_mid(2), RPAN_L, row_mid(2))

    label(sh, RPAN_L, row_top(3) + 0.05, RPAN_W, 1.35,
          'Why the stop-gradient matters\n'
          'Without it, the pathology classification loss could reshape the '
          'lesion decoder to produce convenient pseudo-lesions, improving the '
          'classifier while degrading the segmentation a clinician sees. '
          'Detaching makes the attention read-only.',
          pt=11, color=MUTED)

    label(sh, RPAN_L, row_top(6) - 0.05, RPAN_W, 1.15,
          'What this buys\n'
          'The lesion-type classifier reads features pooled under the model\'s '
          'own predicted lesion map instead of a single global average. Cost: '
          '4,608 extra parameters, one widened linear layer, and no additional '
          'supervision.',
          pt=11, color=MUTED)

    # legend
    ly = row_top(8) + 0.14
    arrow(sh, RPAN_L, ly, RPAN_L + 0.45, ly, head=False)
    label(sh, RPAN_L + 0.58, ly - 0.14, 1.9, 0.28, 'Gradient flows', pt=10,
          color=MUTED)
    arrow(sh, RPAN_L + 2.60, ly, RPAN_L + 3.05, ly, color=GREEN, head=False)
    label(sh, RPAN_L + 3.18, ly - 0.14, 2.2, 0.28, 'Detached (no gradient)',
          pt=10, color=MUTED)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == '__main__':
    path = build()
    print('Wrote', path)
