"""Builds the visual qualitative-results document: original image, then Phase 4, 5 and 6.

Each selected test case gets one full-width figure with four panels in a row, so the images
stay large. Every model panel draws the expert annotation and the model's own prediction on
the same ultrasound, which is the only way to see overlap and error at a glance.

Phase 4 and Phase 5 run on the unenhanced input; Phase 6 runs on images enhanced with the
parameters its optimizer selected, so each panel shows the input that model actually saw.

Usage:  python scripts/14_qualitative_showcase.py --cases 5
"""
import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import matplotlib
matplotlib.use('Agg')
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
import numpy as np
import torch
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from src.config import CHECKPOINTS_DIR, FIGURES_DIR, PROJECT_ROOT, RESULTS_DIR
from src.data import URICADSMultiTaskDataset, load_manifest
from src.engine import get_device, set_seed
from src.enhanced_cache import build_enhanced_splits
from src.metrics import best_threshold, dice_per_sample
from src.model import MultiTaskNetV2

OUT = os.path.join(PROJECT_ROOT, 'URI-CADS_Qualitative_Results.docx')
SHOWCASE_DIR = os.path.join(FIGURES_DIR, 'showcase')
NAVY, INK, SLATE = '21295C', '1A2233', '5B6B79'

GT_KIDNEY, PRED_KIDNEY = '#3ddc84', '#4fc3f7'
GT_LESION, PRED_LESION = '#ffd54f', '#ff5252'

PHASES = [
    ('phase4', 'Phase 4', 'phase4_multitask_local.pt', 'plain',
     'Five-head multi-task scaffold, equal loss weights, unenhanced input.'),
    ('phase5', 'Phase 5', 'phase5_anatomical_local.pt', 'plain',
     'Adds the anatomical consistency constraint (absolute-mass form).'),
    ('phase6', 'Phase 6', 'phase6_optimized_COA_local.pt', 'enhanced',
     'Loss weights and image-enhancement parameters selected by the Coati optimizer.'),
]


# ---------------------------------------------------------------- inference
def load_model(filename, device):
    model = MultiTaskNetV2().to(device)
    state = torch.load(os.path.join(CHECKPOINTS_DIR, filename), map_location=device)
    model.load_state_dict(state)
    model.eval()
    return model


@torch.no_grad()
def predict(model, sample, device):
    image = sample['image'].unsqueeze(0).to(device)
    with torch.autocast('cuda', enabled=device.type == 'cuda'):
        out = model(image)
    kidney_logits = out['kidney_seg'].float()
    lesion_logits = out['lesion_seg'].float()
    k_gt = sample['kidney_mask'].unsqueeze(0).to(device)
    l_gt = sample['lesion_mask'].unsqueeze(0).to(device)
    local_prob = torch.sigmoid(out['local_pathology'].float())[0].cpu().numpy()
    return {
        'kidney': (torch.sigmoid(kidney_logits)[0] > 0.5).cpu().numpy(),
        'lesion': (torch.sigmoid(lesion_logits)[0] > 0.5).cpu().numpy(),
        'kidney_dice': float(dice_per_sample(kidney_logits, k_gt)[0]),
        'lesion_dice': float(dice_per_sample(lesion_logits, l_gt)[0]),
        'diagnosis_prob': float(torch.sigmoid(out['diagnosis'].float())[0].cpu()),
        'local_prob': local_prob,
    }


@torch.no_grad()
def tuned_threshold(model, dataset, device):
    """The operating point each phase would actually use: chosen on validation, never on test."""
    probs, targets = [], []
    for idx in range(len(dataset)):
        sample = dataset[idx]
        image = sample['image'].unsqueeze(0).to(device)
        with torch.autocast('cuda', enabled=device.type == 'cuda'):
            out = model(image)
        probs.append(float(torch.sigmoid(out['diagnosis'].float())[0].cpu()))
        targets.append(int(sample['global_diagnosis'].item()))
    return best_threshold(targets, probs, criterion='balanced_accuracy')


def to_display(tensor):
    """Undo the (x - 0.5) / 0.5 normalization the dataset applies."""
    return (tensor.permute(1, 2, 0).numpy() * 0.5 + 0.5).clip(0, 1)


def scan(models, datasets, device):
    """Per-case predictions and Dice for every phase, over the whole test split."""
    plain = datasets['plain']
    records = []
    for idx in range(len(plain)):
        sample = plain[idx]
        if sample['has_lesion'].item() <= 0:
            continue                      # an empty lesion panel shows nothing useful
        row = {'idx': idx, 'case_id': sample['case_id'],
               'label': int(sample['global_diagnosis'].item())}
        for key, _, _, source, _ in PHASES:
            s = sample if source == 'plain' else datasets[source][idx]
            row[key] = predict(models[key], s, device)
        row['score'] = np.mean([(row[k]['kidney_dice'] + row[k]['lesion_dice']) / 2
                                for k, *_ in PHASES])
        records.append(row)
        if (idx + 1) % 50 == 0:
            print(f'  scanned {idx + 1}/{len(plain)}', flush=True)
    return records


# ---------------------------------------------------------------- rendering
def overlay(ax, base, gt_kidney, gt_lesion, pred):
    ax.imshow(base, cmap='gray' if base.ndim == 2 else None)
    if pred is not None:
        ax.imshow(np.ma.masked_where(pred['kidney'] == 0, pred['kidney']),
                  cmap=matplotlib.colors.ListedColormap([PRED_KIDNEY]), alpha=0.22,
                  vmin=0, vmax=1)
        ax.imshow(np.ma.masked_where(pred['lesion'] == 0, pred['lesion']),
                  cmap=matplotlib.colors.ListedColormap([PRED_LESION]), alpha=0.30,
                  vmin=0, vmax=1)
        for mask, color in ((pred['kidney'], PRED_KIDNEY), (pred['lesion'], PRED_LESION)):
            if mask.any():
                ax.contour(mask.astype(float), levels=[0.5], colors=[color], linewidths=1.9)
        for mask, color in ((gt_kidney, GT_KIDNEY), (gt_lesion, GT_LESION)):
            if mask.any():
                ax.contour(mask, levels=[0.5], colors=[color], linewidths=1.6,
                           linestyles='dashed')
    ax.set_xticks([])
    ax.set_yticks([])
    for spine in ax.spines.values():
        spine.set_visible(False)


def diagnosis_line(pred, thresh, truth):
    """The classification output, stated as the model's actual decision at its own threshold.

    Returned as two short lines rather than one: a single line long enough to hold the verdict,
    the probability and the threshold runs past the panel and collides with its neighbour.
    """
    called = 1 if pred['diagnosis_prob'] >= thresh else 0
    word = 'PATHOLOGICAL' if called else 'HEALTHY'
    mark = 'correct' if called == truth else 'INCORRECT'
    return (f'Diagnosis: {word} - {mark}\n'
            f'p = {pred["diagnosis_prob"]:.3f}  (threshold {thresh:.2f})'), mark == 'correct'


def render_case(record, datasets, thresholds, out_path):
    plain_sample = datasets['plain'][record['idx']]
    gt_kidney = plain_sample['kidney_mask'].numpy()
    gt_lesion = plain_sample['lesion_mask'].numpy()
    truth = record['label']

    fig, axes = plt.subplots(1, 4, figsize=(17.6, 5.6), dpi=150)
    overlay(axes[0], to_display(plain_sample['image']), gt_kidney, gt_lesion, None)
    axes[0].set_title(f'Original ultrasound - case {record["case_id"]}\n\n'
                      f'Reference diagnosis:\n'
                      f'{"PATHOLOGICAL" if truth else "HEALTHY"}',
                      fontsize=10, fontweight='bold', color='#21295C')

    for ax, (key, name, _, source, _) in zip(axes[1:], PHASES):
        sample = plain_sample if source == 'plain' else datasets[source][record['idx']]
        pred = record[key]
        overlay(ax, to_display(sample['image']), gt_kidney, gt_lesion, pred)
        lines, ok = diagnosis_line(pred, thresholds[key], truth)
        ax.set_title(f'{name} prediction\n{lines}\n'
                     f'kidney Dice {pred["kidney_dice"]:.3f}   ·   '
                     f'lesion Dice {pred["lesion_dice"]:.3f}',
                     fontsize=10, fontweight='bold',
                     color='#2E7D5B' if ok else '#B5483B')

    handles = [
        mpatches.Patch(facecolor='none', edgecolor=GT_KIDNEY, linestyle='--', linewidth=1.8,
                       label='Kidney - expert annotation'),
        mpatches.Patch(facecolor=PRED_KIDNEY, edgecolor=PRED_KIDNEY, alpha=0.6,
                       label='Kidney - model prediction'),
        mpatches.Patch(facecolor='none', edgecolor=GT_LESION, linestyle='--', linewidth=1.8,
                       label='Lesion - expert annotation (bounding box)'),
        mpatches.Patch(facecolor=PRED_LESION, edgecolor=PRED_LESION, alpha=0.6,
                       label='Lesion - model prediction'),
    ]
    fig.legend(handles=handles, loc='lower center', ncol=4, frameon=False, fontsize=10,
               bbox_to_anchor=(0.5, -0.035))
    fig.tight_layout()
    fig.savefig(out_path, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return out_path


# ---------------------------------------------------------------- doc helpers
def p(doc, text, *, italic=False, size=10.5, color=INK, space=8, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space)
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return para


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, head in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(head)
        r.bold = True
        r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            cells[i].paragraphs[0].add_run(str(val)).font.size = Pt(9)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Inches(w)
    return t


def dataset_metrics():
    out = {}
    for key, _, _, _, _ in PHASES:
        name = 'phase6_final' if key == 'phase6' else key
        with open(os.path.join(RESULTS_DIR, f'{name}_results.json'), encoding='utf-8') as fh:
            out[key] = json.load(fh)['test_metrics']
    return out


# ---------------------------------------------------------------- document
def build_document(selected, figures, overall, thresholds):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for attr in ('left_margin', 'right_margin'):
        setattr(section, attr, Inches(0.6))
    section.top_margin = section.bottom_margin = Inches(0.55)

    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    for name, size in (('Title', 24), ('Heading 1', 15), ('Heading 2', 12)):
        s = doc.styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(NAVY)
        s.font.bold = True

    doc.add_heading('Qualitative Results', level=0)
    p(doc, 'Diagnosis and segmentation across Phase 4, Phase 5 and Phase 6 on held-out '
           'test images', italic=True, size=12, color=SLATE)

    doc.add_heading('How to read these figures', level=1)
    p(doc, 'Each figure shows one test case. The first panel is the original ultrasound as '
           'acquired, with the reference diagnosis stated beneath it. The next three panels '
           'show the same case processed by successive versions of the model, each drawn on the '
           'input that version actually receives - Phase 4 and Phase 5 work from the unenhanced '
           'image, while Phase 6 uses the enhancement parameters its optimizer selected.')
    p(doc, 'Every model panel reports both of the model\'s output types. The heading gives the '
           'classification result: the diagnosis the model returns, the probability behind it, '
           'the decision threshold in use, and whether that call matches the reference. The '
           'heading is green when the call is correct and red when it is not. The image below '
           'shows the segmentation result.')
    p(doc, 'Dashed outlines are the expert annotation. Filled regions with a solid outline are '
           'the model prediction. Green and blue mark the kidney; yellow and red mark the '
           'lesion. Where a filled region sits inside its dashed outline, the model and the '
           'expert agree.')
    p(doc, f'Thresholds are not left at 0.5. Each phase\'s operating point is chosen on the '
           f'validation split to maximise balanced accuracy and then applied unchanged to test '
           f'- '
           + ', '.join(f'{name} {thresholds[key]:.2f}' for key, name, _, _, _ in PHASES)
           + '. Tuning a threshold on test data would inflate every number reported here.',
      italic=True, color=SLATE)
    p(doc, 'Lesion ground truth is supplied by the dataset as bounding boxes rather than traced '
           'contours, so the dashed yellow outlines are rectangular by construction. The model '
           'predicts a fitted region inside that box, which is clinically more useful but scores '
           'lower against a rectangular target - lesion Dice is therefore pessimistic here.',
      italic=True, color=SLATE)

    doc.add_heading('Test-set performance of each phase', level=1)
    rows = []
    for key, name, _, _, desc in PHASES:
        m = overall[key]
        rows.append([name, f'{m["kidney_dice"]:.4f}', f'{m["lesion_dice_positive_only"]:.4f}',
                     f'{m["diagnosis_auc"]:.4f}',
                     f'{m["consistency_violation_rate"]:.4f}', desc])
    table(doc, ['Phase', 'Kidney Dice', 'Lesion Dice\n(lesion-bearing)', 'Diagnosis AUC',
                'Violation rate', 'Configuration'],
          rows, widths=[0.75, 0.95, 1.05, 0.95, 0.95, 5.4])
    p(doc, '', space=2)
    p(doc, f'Figures computed over all {overall["phase4"]["n_samples"]} held-out test images. '
           f'Kidney segmentation improves steadily across the three phases, and Phase 6 reduces '
           f'the anatomical violation rate roughly fivefold. Lesion Dice does not follow the '
           f'same trend: Phase 4 scores highest, and the optimizer in Phase 6 traded lesion '
           f'overlap for kidney accuracy and anatomical plausibility.', italic=True,
      color=SLATE)

    doc.add_heading('Case-by-case results', level=1)
    p(doc, 'The cases below are drawn from the strongest-performing held-out examples. Every '
           'number printed on a panel is that case\'s own Dice against the expert annotation; '
           'the table above gives the corresponding dataset-wide averages for context.')

    for rank, (record, fig_path) in enumerate(zip(selected, figures), start=1):
        doc.add_heading(f'Case {record["case_id"]}', level=2)
        doc.add_picture(fig_path, width=Inches(9.9))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        label = 'pathological' if record['label'] == 1 else 'healthy'
        bits = '   ·   '.join(
            f'{name} p={record[key]["diagnosis_prob"]:.3f}, kidney '
            f'{record[key]["kidney_dice"]:.3f}, lesion {record[key]["lesion_dice"]:.3f}'
            for key, name, _, _, _ in PHASES)
        p(doc, f'Figure {rank}. Case {record["case_id"]} (reference diagnosis: {label}). {bits}',
          italic=True, size=9, color=SLATE, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()
    doc.add_heading('Notes', level=1)
    for text in [
        'All predictions come from the saved checkpoints of each phase, evaluated at 512 × 512 '
        'on the held-out test split. No model was retrained for this document.',
        'Masks are thresholded at 0.5 on the predicted probability, the same threshold used for '
        'the reported Dice scores.',
        'Lesion ground truth is a union of annotated bounding boxes, so a perfect prediction '
        'cannot reach Dice 1.0 unless the lesion is itself rectangular.',
        'Phase 6 panels appear brighter and smoother than Phase 4 and Phase 5 because its '
        'optimizer selected stronger contrast-limited histogram equalisation and denoising; '
        'this is part of the configuration under test, not a display difference.',
        'The diagnosis shown is the binary healthy-versus-pathological output. The model also '
        'predicts which pathology types are present, as nine independent image-level labels, '
        'but those are not drawn here because the lesion mask is a single untyped channel - '
        'the model states which types occur somewhere in the image, not which type each '
        'individual lesion is.',
    ]:
        para = doc.add_paragraph(style='List Bullet')
        para.paragraph_format.space_after = Pt(5)
        para.add_run(text).font.size = Pt(10.5)

    doc.save(OUT)
    return OUT


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--cases', type=int, default=5)
    args = ap.parse_args()

    os.makedirs(SHOWCASE_DIR, exist_ok=True)
    device = get_device()
    set_seed()
    manifest = load_manifest()

    with open(os.path.join(RESULTS_DIR, 'phase6_final_results.json'), encoding='utf-8') as fh:
        enhancement = json.load(fh)['enhancement']

    print('Building test splits...', flush=True)
    datasets = {'plain': URICADSMultiTaskDataset(manifest[manifest['split'] == 'test'])}
    val = {'plain': URICADSMultiTaskDataset(manifest[manifest['split'] == 'val'])}
    print('Precomputing Phase 6 enhanced images...', flush=True)
    enh = build_enhanced_splits(manifest, enhancement, splits=('val', 'test'), quiet=True)
    datasets['enhanced'], val['enhanced'] = enh['test'], enh['val']

    models = {key: load_model(ckpt, device) for key, _, ckpt, _, _ in PHASES}

    print('Tuning each phase\'s diagnosis threshold on validation...', flush=True)
    thresholds = {}
    for key, name, _, source, _ in PHASES:
        thresholds[key] = tuned_threshold(models[key], val[source], device)
        print(f'  {name}: threshold {thresholds[key]:.3f}', flush=True)

    print(f'Scanning {len(datasets["plain"])} test images across 3 phases...', flush=True)
    records = scan(models, datasets, device)
    records.sort(key=lambda r: r['score'], reverse=True)
    selected = records[:args.cases]

    figures = []
    for record in selected:
        path = os.path.join(SHOWCASE_DIR, f'case_{record["case_id"]}.png')
        print(f'  rendering case {record["case_id"]} (score {record["score"]:.3f})', flush=True)
        figures.append(render_case(record, datasets, thresholds, path))

    out = build_document(selected, figures, dataset_metrics(), thresholds)
    print('Wrote', out)


if __name__ == '__main__':
    main()
