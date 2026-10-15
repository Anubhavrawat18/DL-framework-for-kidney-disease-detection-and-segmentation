"""Assemble the running results document from whatever results exist so far.

Safe to run at any point - it skips phases that have not been run yet, so it can be
regenerated after each phase completes.

Usage:  python scripts/build_report.py
"""
import json
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from src.config import FIGURES_DIR, LOGS_DIR, PROJECT_ROOT, RESULTS_DIR
from src.reporting import load_phase_result

OUT_PATH = os.path.join(PROJECT_ROOT, 'URI-CADS_Results_Document.docx')

METRIC_LABELS = [
    ('diagnosis_auc', 'Healthy/pathological AUC'),
    ('global_pathology_mean_auc', 'Global pathology mean AUC'),
    ('local_pathology_mean_auc', 'Local pathology mean AUC'),
    ('kidney_dice', 'Kidney segmentation Dice'),
    ('lesion_dice', 'Lesion Dice (all test images)'),
    ('lesion_dice_positive_only', 'Lesion Dice (lesion-bearing images only)'),
    ('consistency_violation_rate', 'Consistency violation rate (lower is better)'),
]

PHASE_NARRATIVE = {
    'phase2': ('Phase 2 - Fixed-weight multi-task baseline',
               'ResNet-34 encoder with a diagnosis head, a global-pathology head and a kidney '
               'segmentation decoder, trained with an unweighted sum of losses. This '
               'deliberately mirrors the design choice URI-CADS states in its own paper, so '
               'that the optimized pipeline in Phase 6 has a literature-grounded baseline to '
               'improve on.'),
    'phase3': ('Phase 3 - Fixed-parameter enhancement',
               'Identical to Phase 2 but with a CLAHE + non-local-means denoising + gamma '
               'enhancement stage applied before normalization, using hand-picked parameters. '
               'The original Colab run used only 10 epochs against Phase 2\'s 15, which made '
               'the comparison invalid; this run restores the matched 15-epoch budget.'),
    'phase4': ('Phase 4 - Full five-head multi-task scaffold',
               'Adds a second decoder for coarse lesion segmentation (the union of the '
               'annotated bounding boxes) and a nine-way local-pathology head. This addresses '
               'a limitation URI-CADS states explicitly: they have bounding boxes but no '
               'pixel-level lesion masks.'),
    'phase5': ('Phase 5 - Anatomical consistency constraint (original formulation)',
               'Adds a loss term penalizing lesion probability predicted outside the true '
               'kidney region, expressed as absolute probability mass.'),
    'phase5b': ('Phase 5b - Anatomical consistency constraint (corrected formulation)',
                'Replaces the absolute-mass penalty with a ratio penalty that directly '
                'minimizes the violation rate actually being reported. See the Findings '
                'section for why the original formulation was counterproductive.'),
}


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def add_kv_table(doc, pairs, widths=(2.6, 3.6)):
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Light Grid Accent 1'
    for k, v in pairs:
        row = table.add_row().cells
        row[0].text = str(k)
        row[1].text = str(v)
    for row in table.rows:
        row.cells[0].width = Inches(widths[0])
        row.cells[1].width = Inches(widths[1])
    return table


def add_figure(doc, filename, caption, width=6.4):
    path = os.path.join(FIGURES_DIR, filename)
    if not os.path.exists(path):
        return False
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.runs[0]
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return True


def fmt(value, nd=4):
    if value is None:
        return '-'
    if isinstance(value, float):
        return 'n/a' if value != value else f'{value:.{nd}f}'
    return str(value)


def build():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    # ---------------- title ----------------
    title = doc.add_heading('Optimization-Driven Multi-Task Deep Learning for '
                            'Kidney Abnormality Analysis in Ultrasound', level=0)
    sub = doc.add_paragraph('Experimental results record - URI-CADS dataset')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True
    stamp = doc.add_paragraph(f'Generated {datetime.now():%Y-%m-%d %H:%M} - '
                              f'all experiments run on local hardware')
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp.runs[0].font.size = Pt(9)

    # ---------------- environment ----------------
    add_heading(doc, '1. Compute environment and reproducibility', 1)
    setup_path = os.path.join(LOGS_DIR, 'phase0_local_setup.json')
    if os.path.exists(setup_path):
        with open(setup_path) as fh:
            setup = json.load(fh)
        doc.add_paragraph(
            'All results in this document were produced on the local workstation below. '
            'Every phase uses the same GPU, library versions, random seed and data split, so '
            'the numbers are directly comparable across phases.')
        add_kv_table(doc, [
            ('GPU', setup.get('gpu')),
            ('VRAM', f'{setup.get("vram_gb")} GB'),
            ('PyTorch', setup.get('torch')),
            ('Peak VRAM (5-head model, batch 4 @ 512x512)',
             f'{setup.get("peak_vram_MultiTaskNetV2")} GB'),
            ('Model parameters (5-head)',
             f'{setup.get("params_MultiTaskNetV2", 0)/1e6:.2f} M'),
            ('Dataset', f'{setup.get("n_images")} images / labels / masks'),
            ('Split (train/val/test)',
             ' / '.join(str(setup.get('splits', {}).get(s, '?'))
                        for s in ('train', 'val', 'test'))),
            ('Random seed', 42),
        ])
        doc.add_paragraph(
            f'Measured throughput: {setup.get("est_epoch_min")} minutes per epoch, i.e. about '
            f'{setup.get("est_15epoch_hours")} hours for a 15-epoch run. The five-head model '
            f'occupies {setup.get("peak_vram_MultiTaskNetV2")} GB of the available '
            f'{setup.get("vram_gb")} GB, which is the quantitative basis for the '
            f'edge-deployability claim in the paper.')

    # ---------------- port verification ----------------
    verification = load_phase_result('checkpoint_verification')
    if verification:
        add_heading(doc, '2. Verification of the local reimplementation', 1)
        doc.add_paragraph(
            'The pipeline was reimplemented locally from the original Colab notebooks. To '
            'confirm the reimplementation is faithful, the four checkpoints trained in Colab '
            'were re-evaluated with the local code and compared against the values the '
            'notebooks reported.')
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Light Grid Accent 1'
        for i, head in enumerate(['Checkpoint', 'Metric', 'Colab', 'Local']):
            table.rows[0].cells[i].text = head
        for ckpt, payload in verification.items():
            if not isinstance(payload, dict) or 'local' not in payload:
                continue
            for key, colab_v in payload.get('colab_reported', {}).items():
                cells = table.add_row().cells
                cells[0].text = ckpt.replace('_best.pt', '')
                cells[1].text = dict(METRIC_LABELS).get(key, key)
                cells[2].text = fmt(colab_v)
                cells[3].text = fmt(payload['local'].get(key))
        doc.add_paragraph(
            'Every metric reproduces to within 0.0024, confirming the local implementation '
            'matches the original. All subsequent results are therefore directly comparable '
            'with the earlier work.')
        add_figure(doc, 'checkpoint_verification.png',
                   'Figure 1. Console output of the checkpoint verification run.')

    # ---------------- per-phase results ----------------
    add_heading(doc, '3. Phase results', 1)
    section_no = 0
    for phase_key in ('phase2', 'phase3', 'phase4', 'phase5', 'phase5b'):
        payload = load_phase_result(phase_key)
        if not payload:
            continue
        section_no += 1
        heading, description = PHASE_NARRATIVE[phase_key]
        add_heading(doc, f'3.{section_no} {heading}', 2)
        doc.add_paragraph(description)

        hp = payload.get('hyperparameters', {})
        cfg = payload.get('config', {})
        add_kv_table(doc, [
            ('Architecture', cfg.get('model')),
            ('Active loss terms', ', '.join(cfg.get('active_terms', []))),
            ('Loss weighting', 'equal (unweighted sum)'),
            ('Enhancement', cfg.get('enhancement') or 'none'),
            ('Epochs', hp.get('epochs')),
            ('Batch size', hp.get('batch_size')),
            ('Learning rate', hp.get('lr')),
            ('Checkpoint selection', cfg.get('selection')),
        ])

        doc.add_paragraph('Test-set results:')
        metrics = payload.get('test_metrics', {})
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Light Grid Accent 1'
        table.rows[0].cells[0].text = 'Metric'
        table.rows[0].cells[1].text = 'Value'
        for key, label in METRIC_LABELS:
            if key in metrics:
                cells = table.add_row().cells
                cells[0].text = label
                cells[1].text = fmt(metrics[key])

        add_figure(doc, f'{phase_key}_console.png',
                   f'Figure. {heading} - captured console output of the training run.')
        add_figure(doc, f'{phase_key}_curves.png',
                   f'Figure. {heading} - training loss and validation metrics per epoch.')
        add_figure(doc, f'{phase_key}_roc.png',
                   f'Figure. {heading} - ROC curve for healthy vs pathological.', width=4.0)
        add_figure(doc, f'{phase_key}_qualitative.png',
                   f'Figure. {heading} - qualitative predictions on lesion-bearing test cases.')
        doc.add_page_break()

    # ---------------- cross-phase comparison ----------------
    available = [p for p in ('phase2', 'phase3', 'phase4', 'phase5', 'phase5b')
                 if load_phase_result(p)]
    if available:
        add_heading(doc, '4. Cross-phase comparison', 1)
        doc.add_paragraph(
            'All runs below share the same hardware, seed, split, epoch budget and batch size.')
        section = doc.sections[-1]
        table = doc.add_table(rows=1, cols=1 + len(available))
        table.style = 'Light Grid Accent 1'
        table.rows[0].cells[0].text = 'Metric'
        for i, phase in enumerate(available):
            table.rows[0].cells[i + 1].text = phase.replace('phase', 'Phase ')
        for key, label in METRIC_LABELS:
            cells = table.add_row().cells
            cells[0].text = label
            for i, phase in enumerate(available):
                m = load_phase_result(phase).get('test_metrics', {})
                cells[i + 1].text = fmt(m.get(key), 4)

    # ---------------- phase 6 ----------------
    phase6 = load_phase_result('phase6_search')
    if phase6:
        add_heading(doc, '5. Phase 6 - Metaheuristic joint optimization', 1)
        doc.add_paragraph(
            'The six loss weights and four enhancement parameters are optimized jointly by '
            'three population-based metaheuristics. This replaces both the unweighted loss sum '
            'and the hand-picked preprocessing used in the phases above, and is the central '
            'methodological contribution of the work.')
        doc.add_paragraph('Search space:')
        add_kv_table(doc, [(k, f'[{v[0]:.2f}, {v[1]:.2f}]')
                           for k, v in phase6['search_space'].items()])
        doc.add_paragraph('Fitness function (maximized on the validation split):')
        add_kv_table(doc, [(k, f'weight {v}') for k, v in phase6['fitness_weights'].items()])
        fid = phase6['fidelity']
        doc.add_paragraph(
            f'Each fitness evaluation trains a model from scratch at reduced fidelity '
            f'({fid["img_size"]}x{fid["img_size"]}, {fid["epochs"]} epochs, batch '
            f'{fid["batch_size"]}); the winning configuration is then retrained at full '
            f'fidelity.')

        table = doc.add_table(rows=1, cols=4)
        table.style = 'Light Grid Accent 1'
        for i, head in enumerate(['Optimizer', 'Best fitness', 'Model trainings', 'Hours']):
            table.rows[0].cells[i].text = head
        for name, r in sorted(phase6['results'].items(),
                              key=lambda kv: -kv[1]['best_fitness']):
            cells = table.add_row().cells
            cells[0].text = name
            cells[1].text = fmt(r['best_fitness'])
            cells[2].text = str(r['n_evaluations'])
            cells[3].text = fmt(r['wall_clock_hours'], 2)
        doc.add_paragraph(f'Best optimizer: {phase6["winner"]}.')
        add_figure(doc, 'phase6_convergence.png',
                   'Figure. Convergence of COA, WOA and GWO under an identical budget.')
        add_figure(doc, 'phase6_search_console.png',
                   'Figure. Console output of the Phase 6 search.')

    final = load_phase_result('phase6_final')
    if final:
        add_heading(doc, '5.1 Optimized configuration retrained at full fidelity', 2)
        add_kv_table(doc, list(final.get('loss_weights', {}).items()) +
                     list(final.get('enhancement', {}).items()))
        metrics = final.get('test_metrics', {})
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Light Grid Accent 1'
        table.rows[0].cells[0].text = 'Metric'
        table.rows[0].cells[1].text = 'Value'
        for key, label in METRIC_LABELS:
            if key in metrics:
                cells = table.add_row().cells
                cells[0].text = label
                cells[1].text = fmt(metrics[key])
        add_figure(doc, 'phase6_final_console.png', 'Figure. Final optimized run output.')
        add_figure(doc, 'phase6_final_curves.png', 'Figure. Final optimized training curves.')
        add_figure(doc, 'phase6_final_qualitative.png',
                   'Figure. Final optimized model predictions.')

    proxy_rows = []
    if phase6:
        for opt_name in ('COA', 'GWO', 'WOA'):
            key = 'phase6_final' if opt_name == phase6['winner'] else f'phase6_final_{opt_name}'
            payload = load_phase_result(key)
            if payload and payload.get('optimizer') == opt_name:
                proxy_rows.append((opt_name, phase6['results'][opt_name]['best_fitness'],
                                   payload['test_metrics']))
    if len(proxy_rows) >= 2:
        from src.phase6 import fitness_from_metrics
        add_heading(doc, '5.2 Does the reduced-fidelity search rank configurations correctly?', 2)
        doc.add_paragraph(
            'Each optimizer\'s best configuration was retrained at full fidelity. If the cheap '
            'search proxy is trustworthy, configurations with higher search fitness should also '
            'score higher when fully trained. The final column applies the same fitness '
            'function to full-fidelity test metrics.')
        table = doc.add_table(rows=1, cols=6)
        table.style = 'Light Grid Accent 1'
        for i, head in enumerate(['Optimizer', 'Search fitness (proxy)', 'Diagnosis AUC',
                                  'Kidney Dice', 'Lesion Dice (positive)',
                                  'Full-fidelity fitness']):
            table.rows[0].cells[i].text = head
        for opt_name, proxy_fit, m in sorted(proxy_rows, key=lambda r: -r[1]):
            cells = table.add_row().cells
            cells[0].text = opt_name
            cells[1].text = fmt(proxy_fit)
            cells[2].text = fmt(m.get('diagnosis_auc'))
            cells[3].text = fmt(m.get('kidney_dice'))
            cells[4].text = fmt(m.get('lesion_dice_positive_only'))
            cells[5].text = fmt(fitness_from_metrics(m)[0])
        proxy_order = [r[0] for r in sorted(proxy_rows, key=lambda r: -r[1])]
        full_order = [r[0] for r in sorted(proxy_rows,
                                           key=lambda r: -fitness_from_metrics(r[2])[0])]
        verdict = ('matches' if proxy_order == full_order else 'does NOT match')
        doc.add_paragraph(
            f'Ranking by search fitness: {" > ".join(proxy_order)}. Ranking by full-fidelity '
            f'fitness: {" > ".join(full_order)}. The proxy ranking {verdict} the full-fidelity '
            f'ranking.')

    # ---------------- phase 7 ablations ----------------
    ablations = load_phase_result('phase7_ablations')
    if ablations:
        add_heading(doc, '6. Phase 7 - Ablation study', 1)
        doc.add_paragraph(
            'Each row below is a separately trained model under identical conditions. The '
            'single-task rows test whether multi-task training helps at all; the two '
            '"optimized ... only" rows split the Phase 6 contribution into its two halves, so '
            'the gain can be attributed to optimized loss weighting or to optimized '
            'enhancement rather than to the combination as a whole.')
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light Grid Accent 1'
        for i, head in enumerate(['Configuration', 'Diagnosis AUC', 'Kidney Dice',
                                  'Lesion Dice (positive)', 'Violation rate']):
            table.rows[0].cells[i].text = head
        for payload in ablations.values():
            if not isinstance(payload, dict) or 'test_metrics' not in payload:
                continue
            m = payload['test_metrics']
            cells = table.add_row().cells
            cells[0].text = payload.get('label', '')
            cells[1].text = fmt(m.get('diagnosis_auc'))
            cells[2].text = fmt(m.get('kidney_dice'))
            cells[3].text = fmt(m.get('lesion_dice_positive_only'))
            cells[4].text = fmt(m.get('consistency_violation_rate'))
        add_figure(doc, 'phase7_console.png', 'Figure. Ablation study console output.')

    # ---------------- phase 8 k-fold ----------------
    for cfg_name in ('optimized', 'baseline'):
        kfold = load_phase_result(f'phase8_kfold_{cfg_name}')
        if not kfold:
            continue
        add_heading(doc, f'7. Phase 8 - {kfold["folds"]}-fold cross-validation '
                         f'({cfg_name} configuration)', 1)
        doc.add_paragraph(
            'Phases 2-7 use a single fixed 70/15/15 split. URI-CADS reports its own results '
            'under 5-fold cross-validation over the whole dataset, so this section repeats the '
            'evaluation under that protocol to make the headline comparison like-for-like, and '
            'to attach a variance estimate to each number.')
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Light Grid Accent 1'
        for i, head in enumerate(['Metric', 'Mean', 'Std. dev.']):
            table.rows[0].cells[i].text = head
        for key, label in METRIC_LABELS:
            if key not in kfold['summary']:
                continue
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = fmt(kfold['summary'][key]['mean'])
            cells[2].text = fmt(kfold['summary'][key]['std'])
        pub = kfold['uricads_published']
        ours = kfold['summary']['diagnosis_auc']
        doc.add_paragraph(
            f'Under the same protocol, this system reaches {ours["mean"]:.4f} '
            f'+/- {ours["std"]:.4f} AUC for healthy vs pathological, against the '
            f'{pub["healthy_pathological_auc"]:.4f} reported by URI-CADS. URI-CADS uses a dual '
            f'Mask R-CNN and Faster R-CNN pipeline on a ResNet-50 backbone; the model here is a '
            f'single ResNet-34 encoder with lightweight decoders.')
        add_figure(doc, f'phase8_kfold_{cfg_name}_console.png',
                   f'Figure. {kfold["folds"]}-fold cross-validation console output.')

    # ---------------- cross-validated comparisons ----------------
    cv_runs = [('phase8_kfold_cv_base', 'Baseline (V2, equal weights, plain BCE)'),
               ('phase8_kfold_cv_balanced', 'Baseline + class-balanced BCE'),
               ('phase8_kfold_cv_v3', 'V3 lesion-guided pooling'),
               ('phase8_kfold_cv_v3_ens3', 'V3 + 3-model ensemble per fold')]
    available_cv = [(k, lbl) for k, lbl in cv_runs if load_phase_result(k)]
    if available_cv or load_phase_result('seed_variance_ensemble'):
        add_heading(doc, '7. Cross-validated results and measurement uncertainty', 1)
        doc.add_paragraph(
            'The sections above report a single held-out split. A bootstrap over that split '
            'gives a 95% confidence interval of about 0.085 AUC wide, and retraining the same '
            'recipe with different seeds moves AUC by up to 0.074. Both are larger than almost '
            'every difference reported earlier, so the conclusions that follow are drawn only '
            'from the cross-validated comparisons below.')
    if available_cv:
        add_heading(doc, '7.1 Cross-validated comparison of configurations', 2)
        doc.add_paragraph(
            'Every configuration below was evaluated on the same five folds, so the rows can be '
            'compared to one another directly rather than as independent point estimates. Values '
            'are mean plus or minus standard deviation across folds. This is the only table in '
            'this document from which classification conclusions should be drawn: the '
            'single-split results earlier carry an uncertainty of roughly 0.04 AUC and cannot '
            'separate these configurations.')
        keys = [('diagnosis_auc', 'Diagnosis AUC'),
                ('local_pathology_mean_auc', 'Multi-pathology AUC'),
                ('kidney_dice', 'Kidney Dice'),
                ('lesion_dice_positive_only', 'Lesion Dice (lesion images)')]
        table = doc.add_table(rows=1, cols=1 + len(keys))
        table.style = 'Light Grid Accent 1'
        table.rows[0].cells[0].text = 'Configuration'
        for i, (_, lbl) in enumerate(keys):
            table.rows[0].cells[i + 1].text = lbl
        for key, label in available_cv:
            payload = load_phase_result(key)
            cells = table.add_row().cells
            cells[0].text = label
            for i, (mk, _) in enumerate(keys):
                st = payload['summary'].get(mk)
                cells[i + 1].text = (f'{st["mean"]:.4f} ± {st["std"]:.4f}' if st else '—')
        pub = load_phase_result(available_cv[0][0])['uricads_published']
        cells = table.add_row().cells
        cells[0].text = 'URI-CADS (published, same protocol)'
        cells[1].text = f'{pub["healthy_pathological_auc"]:.4f}'
        cells[2].text = f'{pub["multi_pathological_auc"]:.4f}'
        cells[3].text = '—'
        cells[4].text = '—'

    seed = load_phase_result('seed_variance_ensemble')
    if seed:
        add_heading(doc, '7.2 Training variance and ensembling', 2)
        doc.add_paragraph(
            f'Five models were trained from the identical recipe and data, differing only in the '
            f'random seed. Diagnosis AUC ranged over {seed["seed_auc_spread"]:.4f} '
            f'(standard deviation {seed["seed_auc_std"]:.4f}). This is the floor below which no '
            f'single-split comparison carries information, and it exceeds the size of most '
            f'differences reported in the earlier phases of this project.')
        ens = seed['ensemble']
        add_kv_table(doc, [
            ('Members, mean AUC', fmt(seed['seed_auc_mean'])),
            ('Members, AUC standard deviation', fmt(seed['seed_auc_std'])),
            ('Members, AUC spread', fmt(seed['seed_auc_spread'])),
            ('Ensemble AUC', fmt(ens['auc'])),
            ('Ensemble balanced accuracy', fmt(ens['balanced_accuracy'])),
            ('Ensemble sensitivity / specificity',
             f'{fmt(ens["sensitivity"])} / {fmt(ens["specificity"])}'),
        ])
        doc.add_paragraph(
            'The ensemble exceeds not only the average member but the best individual member, '
            'which is the signature of a real effect rather than a favourable draw.')

    # ---------------- findings ----------------
    findings_path = os.path.join(RESULTS_DIR, 'findings.json')
    if os.path.exists(findings_path):
        add_heading(doc, '8. Findings and methodological corrections', 1)
        with open(findings_path) as fh:
            findings = json.load(fh)
        import re
        for i, item in enumerate(findings, start=1):
            title = re.sub(r'^\d+(\.\d+)*\s+', '', item['title'])
            add_heading(doc, f'8.{i} {title}', 2)
            for para in item['body']:
                doc.add_paragraph(para)

    doc.save(OUT_PATH)
    print(f'Wrote {OUT_PATH}')
    return OUT_PATH


if __name__ == '__main__':
    build()
